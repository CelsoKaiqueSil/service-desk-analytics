"""Build Service Desk KPIs and an executive Excel workbook."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

DATE_COLUMNS = [
    "opened_at",
    "first_response_at",
    "resolved_at",
    "closed_at",
    "sla_due_at",
]
REQUIRED_COLUMNS = {
    "ticket_id",
    "opened_at",
    "first_response_at",
    "resolved_at",
    "closed_at",
    "sla_due_at",
    "category",
    "urgency",
    "owner",
    "status",
}


def load_tickets(path: str | Path, *, start: str | None = None, end: str | None = None) -> pd.DataFrame:
    frame = pd.read_csv(path, dtype={"ticket_id": "string"})
    missing = REQUIRED_COLUMNS.difference(frame.columns)
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
    frame = frame.copy()
    input_rows = len(frame)
    frame['ticket_id'] = frame['ticket_id'].str.strip()
    reasons = pd.Series('', index=frame.index)
    def flag(mask, reason):
        reasons.loc[mask] += reason + '; '
    flag(frame['ticket_id'].isna() | frame['ticket_id'].eq(''), 'missing ticket_id')
    for column in DATE_COLUMNS:
        original = frame[column]
        frame[column] = pd.to_datetime(original, errors="coerce", utc=True)
        flag(original.notna() & original.astype(str).str.strip().ne('') & frame[column].isna(), 'invalid ' + column)
    flag(frame['opened_at'].isna(), 'missing opened_at')
    for column in DATE_COLUMNS[1:]:
        flag(frame[column].notna() & frame[column].lt(frame['opened_at']), column + ' before opening')
    flag(frame['closed_at'].notna() & frame['resolved_at'].notna() & frame['closed_at'].lt(frame['resolved_at']), 'closed before resolution')
    rejected = frame.loc[reasons.ne('')].copy()
    rejected['rejection_reason'] = reasons.loc[reasons.ne('')]
    frame = frame.loc[reasons.eq('')].copy()
    before = len(frame)
    frame = frame.drop_duplicates()
    duplicates = before - len(frame)
    if frame['ticket_id'].duplicated().any():
        raise ValueError('Conflicting versions of ticket_id: reconcile source versions before calculating KPIs.')
    lower = pd.Timestamp(start, tz='UTC') if start else None
    upper = pd.Timestamp(end, tz='UTC') + pd.Timedelta(days=1) if end else None
    if lower is not None and upper is not None and lower >= upper:
        raise ValueError('Start must not be after end.')
    if lower is not None:
        frame = frame.loc[frame['opened_at'].ge(lower)].copy()
    if upper is not None:
        frame = frame.loc[frame['opened_at'].lt(upper)].copy()
        for column in ['first_response_at', 'resolved_at', 'closed_at']:
            frame.loc[frame[column].ge(upper), column] = pd.NaT
    frame["response_hours"] = (
        frame["first_response_at"] - frame["opened_at"]
    ).dt.total_seconds() / 3600
    frame["solution_hours"] = (
        frame["resolved_at"] - frame["opened_at"]
    ).dt.total_seconds() / 3600
    eligible = frame['resolved_at'].notna() & frame['sla_due_at'].notna()
    frame['sla_met'] = pd.Series(pd.NA, index=frame.index, dtype='boolean')
    frame.loc[eligible, 'sla_met'] = frame.loc[eligible, 'resolved_at'].le(frame.loc[eligible, 'sla_due_at'])
    frame.attrs['quality'] = {'input_rows': input_rows, 'rejected_rows': len(rejected), 'duplicates_removed': duplicates,
        'selected_rows': len(frame), 'start': start or 'source minimum', 'end': end or 'source snapshot', 'timezone': 'UTC'}
    frame.attrs['rejected'] = rejected
    return frame


def _distribution(series: pd.Series, prefix: str) -> dict[str, float]:
    values = series.dropna()
    if values.empty:
        return {
            f"{prefix}_mean_hours": float("nan"),
            f"{prefix}_median_hours": float("nan"),
            f"{prefix}_p90_hours": float("nan"),
        }
    return {
        f"{prefix}_mean_hours": float(values.mean()),
        f"{prefix}_median_hours": float(values.median()),
        f"{prefix}_p90_hours": float(values.quantile(0.90)),
    }


def calculate_kpis(frame: pd.DataFrame) -> dict[str, float | int]:
    resolved = frame["resolved_at"].notna()
    closed = frame["closed_at"].notna()
    backlog = frame["status"].str.strip().str.casefold().isin({"aberto", "em andamento"}) & ~resolved & ~closed
    # For an historical opening cohort, current status must not hide work resolved later.
    if frame.attrs.get('quality', {}).get('end') not in {None, 'source snapshot'}:
        backlog = ~resolved & ~closed & ~frame['status'].str.strip().str.casefold().isin({'cancelado', 'cancelled'})
    eligible_sla = resolved & frame["sla_due_at"].notna()
    sla_rate = float(frame.loc[eligible_sla, "sla_met"].mean() * 100) if eligible_sla.any() else float("nan")

    metrics: dict[str, float | int] = {
        "opened": int(len(frame)),
        "resolved": int(resolved.sum()),
        "closed": int(closed.sum()),
        "backlog": int(backlog.sum()),
        "sla_met_percent": sla_rate,
    }
    metrics.update(_distribution(frame["response_hours"], "response"))
    metrics.update(_distribution(frame["solution_hours"], "solution"))
    return metrics


def _segment(frame: pd.DataFrame, column: str) -> pd.DataFrame:
    return (
        frame.groupby(column, dropna=False)
        .agg(
            tickets=("ticket_id", "size"),
            resolved=("resolved_at", lambda values: int(values.notna().sum())),
            response_median_hours=("response_hours", "median"),
            solution_p90_hours=("solution_hours", lambda values: values.dropna().quantile(0.90)),
        )
        .reset_index()
        .sort_values("tickets", ascending=False)
    )


def _excel_safe(frame: pd.DataFrame) -> pd.DataFrame:
    """Return a copy whose datetime columns are compatible with Excel."""
    result = frame.copy()
    for column in result.columns:
        if isinstance(result[column].dtype, pd.DatetimeTZDtype):
            result[column] = result[column].dt.tz_convert("UTC").dt.tz_localize(None)
    return result


def build_report(input_path: str | Path, output_path: str | Path, *, start: str | None = None, end: str | None = None) -> Path:
    frame = load_tickets(input_path, start=start, end=end)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    dashboard = pd.DataFrame(
        [{"indicator": key, "value": value} for key, value in calculate_kpis(frame).items()]
    )
    methodology = pd.DataFrame(
        [
            {"indicator": "backlog", "definition": "Estoque da coorte de abertura selecionada, nao o backlog global. Em recorte historico, usa datas conhecidas ate o fim do dia UTC; cancelados nao entram."},
            {"indicator": "response_hours", "definition": "Primeira resposta menos abertura, em horas."},
            {"indicator": "solution_hours", "definition": "Resolução menos abertura, em horas."},
            {"indicator": "sla_met", "definition": "Resolvido até o prazo, somente quando ambas as datas existem."},
            {"indicator": "p90", "definition": "Valor abaixo do qual ficam 90% dos tempos observados."},
            {"indicator": "period", "definition": "Inicio e fim inclusivos por data de abertura, UTC. Eventos posteriores ao fim nao entram. Sem periodo, usa o snapshot completo de entrada."},
            {"indicator": "deduplication", "definition": "Linhas identicas removidas; IDs conflitantes bloqueiam a geracao. Datas invalidas ou invertidas ficam em Qualidade."},
        ]
    )

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        dashboard.to_excel(writer, sheet_name="Dashboard", index=False)
        _excel_safe(frame).to_excel(writer, sheet_name="Tickets", index=False)
        _segment(frame, "category").to_excel(writer, sheet_name="Categorias", index=False)
        _segment(frame, "owner").to_excel(writer, sheet_name="Responsáveis", index=False)
        methodology.to_excel(writer, sheet_name="Metodologia", index=False)
        quality = pd.DataFrame([{'control': k, 'value': v} for k, v in frame.attrs['quality'].items()])
        quality.to_excel(writer, sheet_name='Qualidade', index=False)
        _excel_safe(frame.attrs['rejected']).to_excel(writer, sheet_name='Qualidade', index=False, startrow=len(quality)+3)
        for sheet in writer.book:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='147E57')
            for column in sheet.columns:
                sheet.column_dimensions[get_column_letter(column[0].column)].width = min(45, max(16, len(str(column[0].value or ''))+3))
        for name, title, value_column in [('Categorias', 'Chamados por categoria', 2), ('Responsáveis', 'P90 de solucao por responsavel (h)', 5)]:
            source = writer.book[name]
            if source.max_row > 1:
                chart = BarChart()
                chart.title = title
                chart.add_data(Reference(source, min_col=value_column, min_row=1, max_row=source.max_row), titles_from_data=True)
                chart.set_categories(Reference(source, min_col=1, min_row=2, max_row=source.max_row))
                chart.width, chart.height = 18, 8
                writer.book['Dashboard'].add_chart(chart, 'D2' if name == 'Categorias' else 'D18')
        writer.book['Dashboard'].sheet_properties.pageSetUpPr.fitToPage = True
        writer.book['Dashboard'].page_setup.orientation = 'landscape'
        writer.book['Dashboard'].page_setup.paperSize = '9'
        writer.book['Dashboard'].page_setup.fitToWidth = 1
        writer.book['Dashboard'].page_setup.fitToHeight = 1
        writer.book['Dashboard'].print_area = 'A1:N34'
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="data/tickets.csv")
    parser.add_argument("--output", default="output/service_desk.xlsx")
    parser.add_argument('--start', help='Opening date YYYY-MM-DD, inclusive, UTC')
    parser.add_argument('--end', help='Opening date YYYY-MM-DD, inclusive, UTC; historical cutoff')
    args = parser.parse_args()
    print(build_report(args.input, args.output, start=args.start, end=args.end))


if __name__ == "__main__":
    main()
